"""Opens the files the BROWSER downloaded and checks them against the database. Usage: verify_ui_files.py <dir>"""
import io, json, subprocess, sys, warnings
import openpyxl, pypdfium2 as pdfium
warnings.filterwarnings("ignore")
d = sys.argv[1]


def pg(sql):
    return subprocess.run(["docker", "exec", "-i", "property_postgres", "psql", "-U", "postgres", "-d", "property_management", "-At"], input=sql, capture_output=True, text=True, encoding="utf-8").stdout.strip()


truth = json.loads(pg("select json_agg(t order by id) from (select p.id, p.title, p.rent_amount, p.is_available, o.full_name owner, t.full_name tenant from properties p left join users o on o.id=p.owner_id left join users t on t.id=p.tenant_id) t"))
out = {}
pdf = pdfium.PdfDocument(open(f"{d}/all_properties.pdf", "rb").read())
text = " ".join(pdf[i].get_textpage().get_text_bounded() for i in range(len(pdf))).replace("\r\n", " ").replace("\n", " ")
out["pdf_opens"] = len(pdf) >= 1
out["pdf_all_properties"] = all(t["title"] in text and f"{int(t['rent_amount']):,}" in text for t in truth)
out["pdf_bandra"] = "Bandra West 2BHK" in text and "65,000" in text and "Vikram Mehta" in text
wb = openpyxl.load_workbook(f"{d}/all_properties.xlsx")
ws = wb["Properties"]
rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
out["xlsx_rows"] = len(rows) == len(truth) and sorted(rows) == [t["id"] for t in truth]
out["xlsx_data"] = all(rows[t["id"]][1] == t["title"] and abs(rows[t["id"]][10] - t["rent_amount"]) < 1e-6 and rows[t["id"]][11] == ("Available" if t["is_available"] else "Occupied") and (rows[t["id"]][12] or None) == t["owner"] and (rows[t["id"]][13] or None) == t["tenant"] for t in truth)
out["xlsx_headings"] = [c.value for c in ws[1]][:3] == ["Property ID", "Property Name", "Address"]
# per-property Bandra West report the browser downloaded
b = pdfium.PdfDocument(open(f"{d}/bandra_report.pdf", "rb").read())
bt = " ".join(b[i].get_textpage().get_text_bounded() for i in range(len(b))).replace("\r\n", " ").replace("\n", " ")
out["bandra_pdf"] = "Bandra West 2BHK" in bt and "301 Turner Road" in bt and "Rs. 65,000" in bt and "No rent payments have been recorded" in bt
bx = openpyxl.load_workbook(f"{d}/bandra_report.xlsx")
out["bandra_xlsx"] = bx.sheetnames == ["Property", "Expenses", "Payments"] and bx["Property"]["B2"].value == "Bandra West 2BHK" and "No expenses" in str(bx["Expenses"]["A4"].value)
print(json.dumps(out))
