"""
Report Service — generates PDF financial reports, rent receipts, and Excel exports.
"""
import io
import logging
from typing import List, Dict, Any
from datetime import datetime

from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.styles import ParagraphStyle as _PS
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer, HRFlowable,
)
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from app.services.export_common import (
    BLUE, LIGHT_BLUE, LIGHT_GRAY, FONT, FONT_BOLD, NumberedCanvas, pdf_markup, pdf_text, set_text,
)

logger = logging.getLogger(__name__)


# ── PDF: Financial report ─────────────────────────────────────────────────────

def _detail_rows(info: Dict) -> List[List[str]]:
    """The property's own details, so a report is meaningful even when it has no transactions yet."""
    def line(label, value):
        return [label, value] if value not in (None, "") else None
    place = ", ".join(x for x in (info.get("city"), info.get("state")) if x)
    rent = info.get("rent")
    rows = [
        line("Property", info.get("title")), line("Address", info.get("address") or "N/A"), line("City / State", place or None),
        line("Type", info.get("type")), line("Monthly rent", f"Rs. {rent:,.0f}" if rent is not None else None),
        line("Status", info.get("status")), line("Owner", info.get("owner")), line("Tenant", info.get("tenant")),
    ]
    return [r for r in rows if r]


def generate_financial_pdf(
    property_info: Dict,
    expenses: List[Dict],
    payments: List[Dict],
    period: str,
    generated_at: datetime = None,
) -> bytes:
    generated_at = generated_at or datetime.now()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=0.6 * inch, bottomMargin=0.7 * inch, title="Financial Report", author="PropAI")
    cell = _PS("cell", fontName=FONT, fontSize=9.5, leading=12)
    cell_b = _PS("cellb", parent=cell, fontName=FONT_BOLD, textColor=BLUE)
    story = []

    story.append(Paragraph("PropAI — AI-Driven Financial Analytics", _PS("AppName", fontName=FONT, fontSize=10, textColor=colors.grey)))
    story.append(Paragraph("Financial Report", _PS("Title", fontName=FONT_BOLD, fontSize=22, leading=26, spaceAfter=4)))
    story.append(HRFlowable(width="100%", thickness=2, color=BLUE, spaceAfter=10))

    info_rows = _detail_rows(property_info) + [["Period", period], ["Generated", generated_at.strftime("%d %b %Y, %H:%M")]]
    info_t = Table([[Paragraph(pdf_markup(k), cell_b), Paragraph(pdf_markup(v), cell)] for k, v in info_rows], colWidths=[1.5 * inch, 5.3 * inch])
    info_t.setStyle(TableStyle([
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("PADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(info_t)
    story.append(Spacer(1, 12))

    total_expenses = sum(e.get("amount", 0) for e in expenses)
    total_rent     = sum(p.get("amount", 0) for p in payments)
    net_income     = total_rent - total_expenses
    summary_data = [
        ["Total Rent Collected", f"Rs. {total_rent:,.2f}"],
        ["Total Expenses",       f"Rs. {total_expenses:,.2f}"],
        ["Net Income",           f"Rs. {net_income:,.2f}"],
        ["No. of Payments",      str(len(payments))],
        ["No. of Expense Records", str(len(expenses))],
    ]
    story.append(Paragraph("Summary", _PS("H2", fontName=FONT_BOLD, fontSize=13, spaceAfter=6)))
    sum_t = Table(summary_data, colWidths=[3.0 * inch, 3.8 * inch])
    sum_t.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (-1, -1), FONT),
        ("FONTNAME",  (0, 0), (0, -1), FONT_BOLD),
        ("FONTSIZE",  (0, 0), (-1, -1), 10),
        ("BACKGROUND",(0, -3), (-1, -3), LIGHT_BLUE),
        ("FONTNAME",  (0, -3), (-1, -3), FONT_BOLD),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("PADDING", (0, 0), (-1, -1), 7),
        ("BOX", (0, 0), (-1, -1), 1, BLUE),
    ]))
    story.append(sum_t)
    story.append(Spacer(1, 16))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.lightgrey, spaceAfter=12))

    empty = _PS("empty", fontName=FONT, fontSize=10, leading=14, textColor=colors.HexColor("#4B5563"))

    # Expenses
    story.append(Paragraph("Expenses", _PS("H2b", fontName=FONT_BOLD, fontSize=14, spaceBefore=6, spaceAfter=6)))
    if not expenses:
        story.append(Paragraph("No expenses have been recorded for this property yet.", empty))
    else:
        rows = [["#", "Category", "Vendor", "Month", "Amount (Rs.)"]]
        total = 0.0
        for i, e in enumerate(expenses, 1):
            rows.append([
                str(i),
                pdf_text((e.get("category") or "").replace("_", " ").title()),
                Paragraph(pdf_markup(e.get("vendor") or "—"), cell),
                pdf_text(e.get("month") or "—"),
                f"Rs. {e.get('amount', 0):,.2f}",
            ])
            total += e.get("amount", 0)
        rows.append(["", "", "", "TOTAL", f"Rs. {total:,.2f}"])
        t = Table(rows, colWidths=[0.4 * inch, 1.6 * inch, 1.8 * inch, 1.0 * inch, 1.4 * inch], repeatRows=1)
        t.setStyle(TableStyle([
            ("FONTNAME",   (0, 0), (-1, -1), FONT),
            ("BACKGROUND", (0, 0), (-1, 0), BLUE),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), FONT_BOLD),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, LIGHT_GRAY]),
            ("BACKGROUND", (0, -1), (-1, -1), LIGHT_BLUE),
            ("FONTNAME",   (3, -1), (-1, -1), FONT_BOLD),
            ("ALIGN",      (4, 0),  (4, -1), "RIGHT"),
            ("VALIGN",     (0, 0),  (-1, -1), "TOP"),
            ("GRID",       (0, 0),  (-1, -1), 0.4, colors.lightgrey),
            ("PADDING",    (0, 0),  (-1, -1), 5),
        ]))
        story.append(t)
    story.append(Spacer(1, 16))

    # Payments
    story.append(Paragraph("Rent Payments", _PS("H2c", fontName=FONT_BOLD, fontSize=14, spaceBefore=6, spaceAfter=6)))
    if not payments:
        story.append(Paragraph("No rent payments have been recorded for this property yet.", empty))
    else:
        prows = [["#", "Month", "Amount (Rs.)", "Status", "Date"]]
        for i, p in enumerate(payments, 1):
            prows.append([
                str(i),
                pdf_text(p.get("month") or "—"),
                f"Rs. {p.get('amount', 0):,.2f}",
                pdf_text((p.get("status") or "").title()),
                pdf_text(str(p.get("payment_date", ""))[:10]),
            ])
        pt = Table(prows, colWidths=[0.4 * inch, 1.0 * inch, 1.4 * inch, 1.2 * inch, 1.8 * inch], repeatRows=1)
        pt.setStyle(TableStyle([
            ("FONTNAME",   (0, 0), (-1, -1), FONT),
            ("BACKGROUND", (0, 0), (-1, 0), BLUE),
            ("TEXTCOLOR",  (0, 0), (-1, 0), colors.white),
            ("FONTNAME",   (0, 0), (-1, 0), FONT_BOLD),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, LIGHT_GRAY]),
            ("ALIGN",      (2, 0), (2, -1), "RIGHT"),
            ("GRID",       (0, 0), (-1, -1), 0.4, colors.lightgrey),
            ("PADDING",    (0, 0), (-1, -1), 5),
        ]))
        story.append(pt)

    class Canvas(NumberedCanvas):
        footer_label = "PropAI · Financial Report"
    doc.build(story, canvasmaker=Canvas)
    return buf.getvalue()


# ── PDF: Rent receipt ─────────────────────────────────────────────────────────

def generate_rent_receipt_pdf(
    tenant_name: str,
    property_address: str,
    amount: float,
    month: str,
    payment_date: str,
) -> bytes:
    buf = io.BytesIO()
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=1 * inch)
    story = []

    story.append(Spacer(1, 20))
    story.append(Paragraph(
        "RENT RECEIPT",
        ParagraphStyle("BigTitle", fontSize=26, fontName="Helvetica-Bold", alignment=TA_CENTER, spaceAfter=20),
    ))
    story.append(HRFlowable(width="80%", thickness=2, color=BLUE, spaceAfter=20))

    data = [
        ["Received From", tenant_name],
        ["Property Address", property_address],
        ["Rent Amount", f"Rs. {amount:,.2f}"],
        ["For the Month of", month],
        ["Payment Date", payment_date],
    ]
    t = Table(data, colWidths=[2.2 * inch, 4 * inch])
    t.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 12),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("PADDING",   (0, 0), (-1, -1), 10),
        ("BOX",       (0, 0), (-1, -1), 1, BLUE),
        ("LINEBEFORE",(1, 0), (1, -1), 1, BLUE),
    ]))
    story.append(t)
    story.append(Spacer(1, 40))
    story.append(Paragraph("Authorised Signature: _______________________", styles["Normal"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "This is a computer-generated receipt.",
        ParagraphStyle("Footer", fontSize=8, textColor=colors.grey, alignment=TA_CENTER),
    ))
    doc.build(story)
    return buf.getvalue()


# ── PDF: Lease Agreement ─────────────────────────────────────────────────────

def generate_lease_pdf(
    tenant_name: str,
    owner_name: str,
    property_address: str,
    property_city: str,
    rent_amount: float,
    start_date: str,
) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=0.8 * inch, bottomMargin=0.8 * inch)
    story = []
    styles = getSampleStyleSheet()

    center_bold = ParagraphStyle("CB", fontSize=14, fontName="Helvetica-Bold", alignment=TA_CENTER, spaceAfter=4)
    center_sm   = ParagraphStyle("CS", fontSize=9,  textColor=colors.grey, alignment=TA_CENTER, spaceAfter=16)
    heading     = ParagraphStyle("H", fontSize=11, fontName="Helvetica-Bold", spaceBefore=12, spaceAfter=4)
    body        = ParagraphStyle("B", fontSize=10, leading=16, spaceAfter=8)

    story.append(Spacer(1, 10))
    story.append(Paragraph("RESIDENTIAL LEASE AGREEMENT", center_bold))
    story.append(Paragraph("Generated by PropAI — AI-Driven Property Management Platform", center_sm))
    story.append(HRFlowable(width="100%", thickness=2, color=BLUE, spaceAfter=16))

    story.append(Paragraph("Parties", heading))
    story.append(Paragraph(
        f"This Lease Agreement is entered into on <b>{start_date}</b> between:", body
    ))
    parties = [
        ["Landlord (Owner)", owner_name],
        ["Tenant",           tenant_name],
        ["Property Address", property_address],
        ["City",             property_city],
    ]
    pt = Table(parties, colWidths=[2.0 * inch, 4.5 * inch])
    pt.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("PADDING",   (0, 0), (-1, -1), 8),
        ("BOX",       (0, 0), (-1, -1), 1, BLUE),
        ("LINEBEFORE",(1, 0), (1, -1), 1, colors.lightgrey),
    ]))
    story.append(pt)

    story.append(Paragraph("Terms", heading))
    terms = [
        ["Commencement Date", start_date],
        ["Lease Duration",    "12 months (renewable)"],
        ["Monthly Rent",      f"Rs. {rent_amount:,.2f}"],
        ["Due Date",          "1st of every month"],
        ["Security Deposit",  f"Rs. {rent_amount * 2:,.2f} (2 months rent)"],
        ["Late Fee",          "Rs. 500 per day after 5-day grace period"],
    ]
    tt = Table(terms, colWidths=[2.0 * inch, 4.5 * inch])
    tt.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (0, -1), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 10),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [colors.white, LIGHT_GRAY]),
        ("PADDING",   (0, 0), (-1, -1), 8),
        ("BOX",       (0, 0), (-1, -1), 1, BLUE),
        ("LINEBEFORE",(1, 0), (1, -1), 1, colors.lightgrey),
    ]))
    story.append(tt)

    story.append(Paragraph("Conditions", heading))
    conditions = [
        "1. Tenant shall pay rent by the 1st of each month via bank transfer or online payment.",
        "2. Tenant shall not sublet or assign the property without written consent of the Landlord.",
        "3. Tenant shall maintain the property in good condition and report maintenance issues promptly.",
        "4. Tenant shall not make structural alterations without prior written approval.",
        "5. Landlord shall be responsible for major structural repairs and maintenance.",
        "6. Either party may terminate this agreement with 30 days written notice.",
        "7. Security deposit will be refunded within 30 days of vacating after deductions for damages.",
        "8. Utilities (electricity, water, gas, internet) are payable by the Tenant directly.",
    ]
    for c in conditions:
        story.append(Paragraph(c, body))

    story.append(Spacer(1, 20))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.lightgrey, spaceAfter=20))

    sig_data = [
        ["Landlord Signature", "", "Tenant Signature", ""],
        [owner_name, "", tenant_name, ""],
        ["Date: ______________", "", "Date: ______________", ""],
    ]
    st = Table(sig_data, colWidths=[2.0 * inch, 0.8 * inch, 2.0 * inch, 0.8 * inch])
    st.setStyle(TableStyle([
        ("FONTNAME",  (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE",  (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("LINEABOVE", (0, 0), (0, 0), 1, BLUE),
        ("LINEABOVE", (2, 0), (2, 0), 1, BLUE),
    ]))
    story.append(st)

    story.append(Spacer(1, 16))
    story.append(Paragraph(
        "This is a computer-generated document produced by PropAI. Please sign and retain a copy.",
        ParagraphStyle("Footer", fontSize=8, textColor=colors.grey, alignment=TA_CENTER),
    ))

    doc.build(story)
    return buf.getvalue()


# ── Excel report ──────────────────────────────────────────────────────────────

def generate_excel_report(expenses: List[Dict], payments: List[Dict], period: str, property_info: Dict = None) -> bytes:
    wb = openpyxl.Workbook()
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(fill_type="solid", fgColor="2563EB")
    note_font = Font(italic=True, color="4B5563")
    thin = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    def header(ws, row, labels):
        for c, label in enumerate(labels, 1):
            cell = ws.cell(row=row, column=c)
            set_text(cell, label)
            cell.font, cell.fill, cell.border = header_font, header_fill, thin
            cell.alignment = Alignment(horizontal="center")

    def widths(ws, first_row, minimum=10, maximum=50):
        for col in ws.columns:
            longest = max((len(str(c.value or "")) for c in col if c.row >= first_row), default=0)
            ws.column_dimensions[col[0].column_letter].width = min(max(longest + 4, minimum), maximum)

    # Property sheet: its own details, present even when there are no transactions
    info_rows = _detail_rows(property_info or {})
    ws0 = wb.active
    ws0.title = "Property"
    header(ws0, 1, ["Field", "Value"])
    for r, (k, v) in enumerate(info_rows + [["Period", period]], 2):
        set_text(ws0.cell(row=r, column=1), k)
        ws0.cell(row=r, column=1).font = Font(bold=True)
        set_text(ws0.cell(row=r, column=2), v)
    widths(ws0, 1)

    # Expenses sheet
    ws1 = wb.create_sheet("Expenses")
    set_text(ws1["A1"], f"Financial Report — {period}")
    ws1["A1"].font = Font(bold=True, size=13)
    header(ws1, 3, ["#", "Category", "Vendor", "Month", "Amount (₹)"])

    total = 0.0
    for i, e in enumerate(expenses, 1):
        r = 3 + i
        ws1.cell(row=r, column=1, value=i)
        set_text(ws1.cell(row=r, column=2), (e.get("category") or "").replace("_", " ").title())
        set_text(ws1.cell(row=r, column=3), e.get("vendor") or "—")
        set_text(ws1.cell(row=r, column=4), e.get("month") or "—")
        ws1.cell(row=r, column=5, value=float(e.get("amount", 0)))
        total += e.get("amount", 0)
    if not expenses:
        set_text(ws1.cell(row=4, column=1), "No expenses have been recorded for this property yet.")
        ws1["A4"].font = note_font
    end = ws1.max_row + 1
    set_text(ws1.cell(row=end, column=4), "TOTAL")
    ws1.cell(row=end, column=5, value=total)
    ws1.cell(row=end, column=4).font = Font(bold=True)
    ws1.cell(row=end, column=5).font = Font(bold=True)
    widths(ws1, 3)

    # Payments sheet
    ws2 = wb.create_sheet("Payments")
    header(ws2, 1, ["Month", "Amount (₹)", "Status", "Payment Date"])
    for r, p in enumerate(payments, 2):
        set_text(ws2.cell(row=r, column=1), p.get("month") or "—")
        ws2.cell(row=r, column=2, value=float(p.get("amount", 0)))
        set_text(ws2.cell(row=r, column=3), (p.get("status") or "").title())
        set_text(ws2.cell(row=r, column=4), str(p.get("payment_date", ""))[:10])
    if not payments:
        set_text(ws2.cell(row=2, column=1), "No rent payments have been recorded for this property yet.")
        ws2["A2"].font = note_font
    widths(ws2, 1)

    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
